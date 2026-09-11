# -*- coding: utf-8 -*-
"""Построение графика санитарного разрыва на следующий месяц.

Опирается на график предыдущего месяца:

  * шаблон цикла (последовательность отметок и интервалы между ними) берётся
    из самих птичников — по завершённым циклам этого цеха; ключ --templates-from
    ограничивает набор графиков, по которым шаблон строится;
  * месяцев-источников можно указать несколько (от раннего к позднему), тогда
    история птичника склеивается и в расчёт берётся последнее событие;
  * незакрытые на конец месяца циклы продолжаются по шаблону;
  * после заселения (ЗС) птица растёт заданное число дней, затем начинается
    новый цикл с освобождения (ОС/ОЧН).

    python3 src/make_next_month.py data/grafik-sanrazryva-2026-08.xlsx \
        data/grafik-sanrazryva-2026-09.xlsx \
        --month 2026-10 --days 40 --out data/grafik-sanrazryva-2026-10.xlsx
"""

import argparse
import calendar
import collections
import os
import re
import sys
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_schedule import SHOP_BLOCKS, HEADER_ROW, COL_FLOOR, FIRST_DAY_COL  # noqa: E402

START_MARK = "ОС"       # с чего начинается цикл
END_MARK = "ЗС"         # чем заканчивается

HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")

# Заливка отметок — как в исходном графике: красным начало санразрыва,
# синим мойка, зелёным подготовка, жёлтым заселение. Остальное без заливки.
MARK_FILLS = [
    ("ОС", PatternFill("solid", fgColor="FFFF0000")),
    ("ПДГ", PatternFill("solid", fgColor="FF92D050")),
    ("ЗС", PatternFill("solid", fgColor="FFFFFF00")),
    ("МС", PatternFill("solid", fgColor="FF4F81BD")),
]


def mark_fill(mark):
    for key, fill in MARK_FILLS:
        if key in mark:
            return fill
    return None
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_month(path, sheet=None):
    """[(цех, этаж, строка, [(дата, отметка)])] — все отметки как есть."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]

    days = {}
    for col in range(FIRST_DAY_COL, FIRST_DAY_COL + 31):
        value = ws.cell(HEADER_ROW, col).value
        if hasattr(value, "date"):
            days[col] = value.date()

    houses = []
    for shop, row_from, row_to in SHOP_BLOCKS:
        for row in range(row_from, row_to + 1):
            floor = ws.cell(row, COL_FLOOR).value
            if floor is None:
                continue
            marks = []
            for col, day in sorted(days.items()):
                value = ws.cell(row, col).value
                if value is not None and str(value).strip():
                    marks.append((day, str(value).strip()))
            houses.append({"shop": shop, "floor": floor, "row": row, "marks": marks})
    return houses, sorted(days.values())


def merge_months(sources, sheet=None):
    """История птичника из нескольких месяцев подряд, отсортированная по датам."""
    merged = {}
    order = []
    for path in sources:
        houses, _ = read_month(path, sheet)
        for house in houses:
            key = (house["shop"], house["floor"])
            if key not in merged:
                merged[key] = {"shop": house["shop"], "floor": house["floor"], "marks": []}
                order.append(key)
            merged[key]["marks"].extend(house["marks"])
    for house in merged.values():
        house["marks"].sort(key=lambda x: x[0])
    return [merged[k] for k in order]


def cycle_template(marks, start_date):
    """[(сдвиг в днях, отметка)] от начала цикла."""
    return [((day - start_date).days, mark) for day, mark in marks]


def duplicate_codes(template):
    """Коды, встречающиеся в цикле больше одного раза (ДЗ идёт дважды штатно)."""
    counter = collections.Counter()
    for _, mark in template:
        counter.update(mark_codes(mark))
    return {code for code, n in counter.items() if n > 1 and code not in ("ДЗ", "Р")}


def build_templates(houses):
    """Цех -> шаблон цикла.

    Берём самый частый завершённый цикл цеха. При равенстве частот
    предпочитаем цикл без повторов операций: редкий вариант с двойной
    обработкой, размноженный на весь цех, завышает расход.
    """
    per_shop = collections.defaultdict(collections.Counter)
    overall = collections.Counter()

    for house in houses:
        marks = house["marks"]
        starts = [i for i, (_, m) in enumerate(marks) if START_MARK in m]
        for i in starts:
            tail = marks[i:]
            ends = [j for j, (_, m) in enumerate(tail) if m == END_MARK]
            if not ends:
                continue
            cycle = tail[:ends[0] + 1]
            template = tuple(cycle_template(cycle, cycle[0][0]))
            per_shop[house["shop"]][template] += 1
            overall[template] += 1

    def best(counter):
        return max(counter.items(),
                   key=lambda kv: (kv[1], -len(duplicate_codes(kv[0]))))[0]

    fallback = best(overall) if overall else None
    chosen = {shop: best(counter) for shop, counter in per_shop.items()}

    warnings = []
    for shop, template in sorted(chosen.items()):
        repeats = duplicate_codes(template)
        if repeats:
            warnings.append(
                "цех {}: в шаблоне цикла операция {} встречается дважды — "
                "так в исходном графике, проверьте, что это не опечатка"
                .format(shop, ", ".join(sorted(repeats))))
        if per_shop[shop][template] == 1:
            warnings.append(
                "цех {}: шаблон построен по единственному завершённому циклу — "
                "остальные выходят за границы присланных месяцев"
                .format(shop))
    return chosen, fallback, warnings


def mark_codes(mark):
    """'ЗО/ДЗ\\ГГ' -> {'ЗО', 'ДЗ', 'ГГ'} — для сравнения отметок между собой."""
    return {p.strip().upper() for p in re.split(r"[\\/|,]+", str(mark)) if p.strip()}


def resume_point(template, done):
    """Индекс шага шаблона, на котором цикл уже остановился.

    Сопоставляем по самим отметкам, а не по номеру дня: реальный график
    бывает сдвинут относительно шаблона, и сопоставление по дням приводило
    к повторной простановке уже выполненных работ через границу месяца.
    """
    index = -1
    position = 0
    for _, mark in done:
        codes = mark_codes(mark)
        for step in range(position, len(template)):
            if mark_codes(template[step][1]) & codes:
                index, position = step, step + 1
                break
    return index


def open_cycle(house):
    """Незакрытый цикл: дата начала и уже проставленные отметки, иначе None."""
    marks = house["marks"]
    starts = [i for i, (_, m) in enumerate(marks) if START_MARK in m]
    if not starts:
        return None
    tail = marks[starts[-1]:]
    if any(m == END_MARK for _, m in tail):
        return None
    return tail[0][0], tail


def last_settlement(house):
    """Дата последнего заселения в месяце."""
    days = [day for day, mark in house["marks"] if mark == END_MARK]
    return max(days) if days else None


def generate(houses, templates, fallback, month_from, month_to, grow_days):
    """Птичник -> {дата: отметка} на новый месяц."""
    plan = {}
    unknown = []

    for house in houses:
        template = templates[0].get(house["shop"], fallback)
        if template is None:
            continue
        cycle_len = template[-1][0]
        marks = {}

        # 1. Продолжаем цикл, не закрытый к концу прошлого месяца.
        settled = last_settlement(house)
        opened = open_cycle(house)
        if opened:
            start, done = opened
            index = resume_point(template, done)
            last_day, last_offset = done[-1][0], (template[index][0] if index >= 0 else 0)
            for offset, mark in template[index + 1:]:
                day = last_day + timedelta(days=offset - last_offset)
                if month_from <= day <= month_to:
                    marks[day] = mark
            settled = last_day + timedelta(days=cycle_len - last_offset)

        if settled is None:
            unknown.append("{} / этаж {}".format(house["shop"], house["floor"]))
            continue

        # 2. Новые циклы: через grow_days после заселения.
        start = settled + timedelta(days=grow_days)
        while start <= month_to:
            for offset, mark in template:
                day = start + timedelta(days=offset)
                if month_from <= day <= month_to:
                    marks[day] = mark
            start = start + timedelta(days=cycle_len + grow_days)

        if marks:
            plan[(house["shop"], house["floor"])] = marks

    return plan, unknown


def write_schedule(source, plan, month_from, month_to, out_path, title):
    """Записываем новый месяц чистой книгой, сохраняя раскладку строк.

    Строки птичников остаются на тех же местах, что и в исходном графике
    (см. SHOP_BLOCKS), поэтому файл разбирается тем же parse_schedule.
    """
    src = openpyxl.load_workbook(source, data_only=True)
    src_ws = src.worksheets[0]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title

    days = [month_from + timedelta(days=i)
            for i in range((month_to - month_from).days + 1)]

    ws["A6"] = ('График  санитарного разрыва  производственных цехов  на   {}  '
                'АО "Усть-Каменогорская птицефабрика"'.format(title))
    ws["A6"].font = Font(bold=True, size=12)

    header = [("A", "Цех"), ("B", "Этаж")]
    for col, text in header:
        cell = ws["{}{}".format(col, HEADER_ROW)]
        cell.value = text
    for i, day in enumerate(days):
        cell = ws.cell(HEADER_ROW, FIRST_DAY_COL + i)
        cell.value = day
        cell.number_format = "DD.MM"

    mirror_floor = FIRST_DAY_COL + len(days)
    ws.cell(HEADER_ROW, mirror_floor).value = "Этаж"
    ws.cell(HEADER_ROW, mirror_floor + 1).value = "Цех"

    for col in range(1, mirror_floor + 2):
        cell = ws.cell(HEADER_ROW, col)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for shop, row_from, row_to in SHOP_BLOCKS:
        for row in range(row_from, row_to + 1):
            floor = src_ws.cell(row, COL_FLOOR).value
            if floor is None:
                continue
            ws.cell(row, 1).value = shop if row == row_from else None
            ws.cell(row, COL_FLOOR).value = floor
            ws.cell(row, mirror_floor).value = floor
            ws.cell(row, mirror_floor + 1).value = shop if row == row_from else None
            for day, mark in plan.get((shop, floor), {}).items():
                cell = ws.cell(row, FIRST_DAY_COL + (day - month_from).days)
                cell.value = mark
                fill = mark_fill(mark)
                if fill:
                    cell.fill = fill
                    if fill.start_color.rgb == "FF4F81BD":
                        cell.font = Font(color="FFFFFFFF")
            for col in range(1, mirror_floor + 2):
                cell = ws.cell(row, col)
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")

    legend = [
        ("ос", "освобождение", "ДС", "дезинсекция", "СО", "септик обработка"),
        ("ОЧ", "очистка", "ГФ", "газация формалином", "ОД", "обработка дорог"),
        ("ОЧН", "очистка ночь", "ГГ", "газация глутаркой", "МС", "мойка сутки"),
        ("П", "побелка день", "мд", "мойка ДЕНЬ", "МН", "мойка ночь"),
        ("ПН", "побелка ночь", "р", "ремонт", "Ш", "ШАХТА"),
        ("ДЗ", "дезинфекция", "ПДГ", "подготовка", "ПР", "приемка птичника"),
        ("ЗО", "засыпка", "ЗС", "заселение", "ЗН", "засыпка в ночь"),
    ]
    first_legend_row = 132
    for i, line in enumerate(legend):
        row = first_legend_row + i * 2
        for j, value in enumerate(line):
            ws.cell(row, FIRST_DAY_COL + j + (j // 2)).value = value

    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 7
    for col in range(FIRST_DAY_COL, mirror_floor + 2):
        ws.column_dimensions[get_column_letter(col)].width = 9
    ws.freeze_panes = "C{}".format(HEADER_ROW + 1)

    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="+",
                    help="графики предыдущих месяцев, от раннего к позднему")
    ap.add_argument("--month", required=True, help="новый месяц в виде ГГГГ-ММ")
    ap.add_argument("--days", type=int, default=40, help="длительность выращивания, дней")
    ap.add_argument("--templates-from", nargs="*", metavar="XLSX",
                    help="графики, по которым строить шаблон цикла (по умолчанию все). "
                         "Указывайте только настоящие графики: если брать шаблон из "
                         "ранее построенных месяцев, редкий вариант цикла "
                         "размножается и подменяет собой обычный")
    ap.add_argument("--sheet", help="лист исходного графика")
    ap.add_argument("--out", required=True)
    ap.add_argument("--title", help="имя листа нового графика")
    args = ap.parse_args(argv)

    year, month = (int(x) for x in args.month.split("-"))
    month_from = date(year, month, 1)
    month_to = date(year, month, calendar.monthrange(year, month)[1])

    houses = merge_months(args.xlsx, args.sheet)
    source = merge_months(args.templates_from, args.sheet) if args.templates_from else houses
    templates = build_templates(source)
    plan, unknown = generate(houses, templates, templates[1],
                             month_from, month_to, args.days)

    title = args.title or "{} {}".format(
        ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
         "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"][month - 1], str(year)[2:])
    write_schedule(args.xlsx[-1], plan, month_from, month_to, args.out, title)

    print("График построен:", args.out)
    print("  период: {} — {}, выращивание {} дней".format(month_from, month_to, args.days))
    print("  птичников с отметками: {} из {}".format(len(plan), len(houses)))
    print("  шаблон цикла (общий): " + "  ".join(
        "+{}:{}".format(o, m) for o, m in templates[1]))
    for shop, template in sorted(templates[0].items()):
        if template != templates[1]:
            print("  шаблон цеха {}: ".format(shop) + "  ".join(
                "+{}:{}".format(o, m) for o, m in template))
    for text in templates[2]:
        print("  ВНИМАНИЕ: {}".format(text))
    for house in unknown:
        print("  ВНИМАНИЕ: нет данных для планирования: {}".format(house))
    return 0


if __name__ == "__main__":
    sys.exit(main())
