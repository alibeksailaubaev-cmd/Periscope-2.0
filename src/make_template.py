# -*- coding: utf-8 -*-
"""Сборка Excel-шаблона «Санитарная программа» для ручного заполнения.

Из графика санитарного разрыва берутся все связки «цех × операция» с
количеством обработок за месяц, из прайса — список средств. Дальше в книге
вручную проставляются площадка, средство, норма расхода и цена; потребность
и сумма считаются формулами.

    python3 src/make_template.py data/grafik-sanrazryva-2026-09.xlsx \
        --price data/dezsredstva.csv --out out/sanitarnaya-programma.xlsx
"""

import argparse
import csv
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_schedule import parse  # noqa: E402

# Порядок операций — как они идут по циклу санации.
CODE_ORDER = ["ОС", "ОЧН", "МС", "ПР", "ДС", "ДЗ", "ГГ", "ГФ", "ОД", "ПДГ"]

LEGEND = [
    ("ОС", "освобождение", "санация 1"),
    ("ОЧН", "очистка ночь", "санация 1"),
    ("МС", "мойка сутки", "санация 2"),
    ("ПР", "приёмка птичника", "санация 3"),
    ("ДС", "дезинсекция", ""),
    ("ДЗ", "дезинфекция", ""),
    ("ГГ", "газация глутаркой", ""),
    ("ГФ", "газация формалином", ""),
    ("ОД", "обработка дорог", ""),
    ("ПДГ", "подготовка", "санация 4"),
]

HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_price_list(path):
    """Прайс -> [(наименование, ед.изм, цена|None, примечание)]"""
    items = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh, delimiter=";"):
            name = (row.get("Наименование") or "").strip()
            if not name:
                continue
            raw_price = (row.get("Цена") or "").strip().replace(" ", "").replace(",", ".")
            items.append((
                name,
                (row.get("Ед.изм") or "").strip(),
                float(raw_price) if raw_price else None,
                (row.get("Примечание") or "").strip(),
            ))
    return items


def style_header(ws, row, last_col):
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_price_sheet(wb, price_items):
    ws = wb.create_sheet("Дезсредства")
    ws.append(["№", "Наименование", "Ед.изм", "Цена за ед.", "Примечание"])
    style_header(ws, 1, 5)
    for i, (name, unit, price, note) in enumerate(price_items, start=1):
        ws.append([i, name, unit, price, note])
        ws.cell(ws.max_row, 4).fill = INPUT_FILL
        ws.cell(ws.max_row, 4).number_format = "# ##0.000"
        for col in range(1, 6):
            ws.cell(ws.max_row, col).border = BORDER
    for col, width in zip("ABCDE", (5, 62, 10, 14, 46)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:E{}".format(ws.max_row)
    return ws


def build_legend_sheet(wb):
    ws = wb.create_sheet("Справочник")
    ws.append(["Код", "Операция", "Этап санации"])
    style_header(ws, 1, 3)
    for code, name, stage in LEGEND:
        ws.append([code, name, stage])
        for col in range(1, 4):
            ws.cell(ws.max_row, col).border = BORDER
    ws.append([])
    ws.append(["Не учитываются: Р (ремонт), ЗН/ЗО (засыпка), ЗС (заселение), "
               "СО (септик-обработка), П/ПН (побелка), Ш (шахта)"])
    for col, width in zip("ABC", (8, 28, 16)):
        ws.column_dimensions[col].width = width
    return ws


def build_program_sheet(wb, counts, price_rows, period):
    ws = wb.create_sheet("Санитарная программа", 0)
    ws["A1"] = "Санитарная программа: назначение дезсредств по площадкам"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Период графика: {} — {}. Жёлтые колонки заполняются вручную, " \
               "серые считаются формулами.".format(period["from"], period["to"])
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A3"] = "Если на одну операцию идёт несколько средств — скопируйте строку " \
               "и укажите в копии другое средство; формулы копируются вместе со строкой."
    ws["A3"].font = Font(italic=True, color="666666")

    header = ["Площадка", "Цех", "Код", "Операция", "Обработок за период",
              "Дезсредство", "Ед.изм", "Норма на 1 обработку", "Потребность",
              "Цена за ед.", "Сумма", "Примечание"]
    ws.append(header)
    head_row = ws.max_row
    style_header(ws, head_row, len(header))

    price_last = len(price_rows) + 1
    names_ref = "Дезсредства!$B$2:$B${}".format(price_last)
    lookup_ref = "Дезсредства!$B$2:$D${}".format(price_last)

    dv = DataValidation(type="list", formula1="={}".format(names_ref), allow_blank=True)
    dv.error = "Выберите средство из списка на листе «Дезсредства»"
    dv.errorTitle = "Нет такого средства"
    ws.add_data_validation(dv)

    for shop, code, name, qty in counts:
        ws.append([None, shop, code, name, qty, None, None, None, None, None, None, None])
        r = ws.max_row
        ws.cell(r, 7).value = '=IFERROR(VLOOKUP(F{r},{ref},2,FALSE),"")'.format(r=r, ref=lookup_ref)
        ws.cell(r, 9).value = '=IF(H{r}="","",E{r}*H{r})'.format(r=r)
        ws.cell(r, 10).value = '=IFERROR(VLOOKUP(F{r},{ref},3,FALSE),"")'.format(r=r, ref=lookup_ref)
        ws.cell(r, 11).value = '=IF(OR(I{r}="",J{r}=""),"",I{r}*J{r})'.format(r=r)
        dv.add(ws.cell(r, 6))
        for col in range(1, len(header) + 1):
            ws.cell(r, col).border = BORDER
        for col in (1, 6, 8, 12):
            ws.cell(r, col).fill = INPUT_FILL
        for col in (7, 9, 10, 11):
            ws.cell(r, col).fill = CALC_FILL
        ws.cell(r, 5).number_format = "# ##0"
        ws.cell(r, 8).number_format = "# ##0.000"
        ws.cell(r, 9).number_format = "# ##0.000"
        ws.cell(r, 10).number_format = "# ##0.000"
        ws.cell(r, 11).number_format = "# ##0.00"

    first, last = head_row + 1, ws.max_row
    ws.append([])
    total = ws.max_row + 1
    ws.cell(total, 4, "ИТОГО").font = Font(bold=True)
    ws.cell(total, 11, "=SUM(K{}:K{})".format(first, last)).font = Font(bold=True)
    ws.cell(total, 11).number_format = "# ##0.00"
    ws.cell(total, 11).fill = CALC_FILL
    ws.cell(total, 11).border = BORDER

    for col, width in zip("ABCDEFGHIJKL",
                          (16, 9, 7, 22, 12, 46, 9, 14, 14, 14, 16, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A{}".format(head_row + 1)
    ws.auto_filter.ref = "A{}:L{}".format(head_row, last)
    return ws


def collect_counts(data):
    """[(цех, код, операция, количество обработок)] в порядке цикла санации."""
    counter = Counter()
    names = {}
    for house in data["houses"]:
        for e in house["events"]:
            counter[(house["shop"], e["code"])] += 1
            names[e["code"]] = e["name"]
    rows = []
    for shop in sorted({k[0] for k in counter}):
        for code in CODE_ORDER:
            qty = counter.get((shop, code))
            if qty:
                rows.append((shop, code, names[code], qty))
    return rows


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="график санитарного разрыва")
    ap.add_argument("--price", default="data/dezsredstva.csv", help="прайс дезсредств (CSV)")
    ap.add_argument("--sheet", help="лист графика (по умолчанию первый)")
    ap.add_argument("--out", default="out/sanitarnaya-programma.xlsx", help="куда сохранить книгу")
    args = ap.parse_args(argv)

    data = parse(args.xlsx, args.sheet)
    counts = collect_counts(data)
    price_items = read_price_list(args.price)

    wb = Workbook()
    wb.remove(wb.active)
    build_program_sheet(wb, counts, price_items, data["period"])
    build_price_sheet(wb, price_items)
    build_legend_sheet(wb)

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(args.out)

    print("Книга собрана:", args.out)
    print("  строк для заполнения: {} (цех × операция)".format(len(counts)))
    print("  средств в справочнике: {}".format(len(price_items)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
