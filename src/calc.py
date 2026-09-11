# -*- coding: utf-8 -*-
"""Расчёт расхода и стоимости дезсредств за месяц.

Связывает три источника:
  * график санитарного разрыва — сколько раз за месяц проводится каждая
    операция в каждом цехе;
  * санитарная программа (нормы) — сколько средства уходит на один птичник;
  * прайс — цена за единицу.

    python3 src/calc.py data/grafik-sanrazryva-2026-09.xlsx \
        --out out/raschet-2026-09.xlsx
"""

import argparse
import csv
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_schedule import parse  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# Средства, у которых норма задана в литрах, а цена в килограммах (или наоборот).
# Такие пары считаются один к одному, а список выводится отдельным замечанием.
LITRE_KILO = {"л", "кг", "л (дм3)"}
UNIT_MISMATCHES = {}


def note_unit_mismatch(agent, norm_unit, price_unit):
    norm_unit = (norm_unit or "").strip().lower()
    price_unit = (price_unit or "").strip().lower()
    if not norm_unit or not price_unit or norm_unit == price_unit:
        return
    if norm_unit in LITRE_KILO and price_unit in LITRE_KILO:
        UNIT_MISMATCHES[agent] = (norm_unit, price_unit)


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=";"))


def to_float(value):
    value = (value or "").strip().replace(" ", "").replace(",", ".")
    return float(value) if value else None


def load_prices(price_path, map_path):
    """Название средства в программе -> (позиция прайса, цена, ед.изм)."""
    prices = {}
    for row in read_csv(price_path):
        name = (row.get("Наименование") or "").strip()
        if name:
            prices[name] = (to_float(row.get("Цена")), (row.get("Ед.изм") or "").strip())
    mapping = {}
    for row in read_csv(map_path):
        src = (row.get("Средство в программе") or "").strip()
        dst = (row.get("Позиция прайса") or "").strip()
        price, unit = prices.get(dst, (None, ""))
        mapping[src] = (dst, price, unit)
    return mapping


def count_operations(schedule):
    """(цех, код, № вхождения) -> сколько раз операция проводится за период.

    Операция может повторяться за цикл (дезинфекция идёт дважды: на день 5 и
    на день 6), поэтому вхождения внутри каждого птичника нумеруются по дате.
    """
    counter = Counter()
    for house in schedule["houses"]:
        seen = Counter()
        for event in sorted(house["events"], key=lambda e: e["date"]):
            seen[event["code"]] += 1
            counter[(house["shop"], event["code"], seen[event["code"]])] += 1
    return counter


def treatments_for(counts, shops, codes, occurrence):
    total = 0
    for shop in shops:
        for code in codes:
            for (c_shop, c_code, c_occ), qty in counts.items():
                if c_shop == shop and c_code == code:
                    if not occurrence or c_occ == occurrence:
                        total += qty
    return total


def build_extra_rows(extra, prices, houses_per_shop):
    """Работы, идущие не по графику санразрыва: дератизация, коврики, барьеры.

    Норма задаётся на птичник, корпус, барьер или сразу на площадку, поэтому
    количество за месяц берётся из самой таблицы: единиц × раз в месяц.
    """
    rows = []
    for norm in extra:
        agent = norm["Средство"].strip()
        position, price, price_unit = prices.get(agent, ("", None, ""))
        per_unit = to_float(norm["Норма"])
        qty = int(float(norm["Единиц"]) * float(norm["Раз в месяц"]))
        month_qty = per_unit * qty if per_unit is not None else None
        total = month_qty * price if month_qty is not None and price is not None else None

        problems = []
        if price is None:
            problems.append("нет цены")
        note_unit_mismatch(agent, norm["Ед.изм"], price_unit)
        # В «Замечания» выносим только то, что требует решения, а не пояснения.
        if "уточнить" in (norm.get("Примечание") or "").lower():
            problems.append(norm["Примечание"])

        shops = [s.strip() for s in norm["Цеха"].split(",") if s.strip()]
        rows.append({
            "Приложение": norm["Приложение"],
            "Процесс": norm["Процесс"],
            "Код": "—",
            "Площадка": norm["Площадка"],
            "Цеха": norm["Цеха"],
            "Птичников на площадке": int(sum(houses_per_shop.get(s, 0) for s in shops))
                                     or sum(houses_per_shop.values()),
            "Средство": agent,
            "Позиция прайса": position,
            "Ед.изм": norm["Ед.изм"],
            "Норма": per_unit,
            "Норма на": norm["Норма на"],
            "Раствор на 1 птичник, л": None,
            "Кол-во за месяц": qty,
            "Расход за месяц": month_qty,
            "Раствор за месяц, л": None,
            "Цена за ед.": price,
            "Сумма": total,
            "Замечания": "; ".join(problems),
        })
    return rows


def build_rows(schedule, norms, prices):
    counts = count_operations(schedule)
    grouped = defaultdict(lambda: {"shops": [], "norm": None})

    houses_per_shop = {}
    for norm in norms:
        key = (norm["Приложение"], norm["Процесс"], norm["Код в графике"],
               norm["Участок"], norm["Средство"], norm.get("Вхождение", ""))
        grouped[key]["shops"].append(norm["Цех"])
        houses_per_shop[norm["Цех"]] = to_float(norm.get("Птичников в цехе")) or 0
        grouped[key]["norm"] = norm

    rows = []
    for (app, process, code, site, agent, _occ), data in grouped.items():
        norm = data["norm"]
        codes = code.split("+")
        occurrence = int(norm["Вхождение"]) if (norm.get("Вхождение") or "").strip() else None
        shops = set(data["shops"])
        treatments = treatments_for(counts, shops, codes, occurrence)
        houses = int(sum(houses_per_shop.get(shop, 0) for shop in shops))
        per_house = to_float(norm["Норма средства на птичник"])
        position, price, price_unit = prices.get(agent, ("", None, ""))
        solution = to_float(norm["Раствор, л"])

        month_qty = per_house * treatments if per_house is not None else None
        total = month_qty * price if month_qty is not None and price is not None else None

        problems = []
        if per_house is None:
            problems.append("нет нормы")
        if price is None:
            problems.append("нет цены")
        if treatments == 0:
            problems.append("нет обработок в графике")
        elif treatments != houses and code != "ГГ+ГФ":
            problems.append("обработок {} при {} птичниках — цикл выходит за месяц"
                            .format(treatments, houses))
        note_unit_mismatch(agent, norm["Ед.изм"], price_unit)

        rows.append({
            "Приложение": app,
            "Процесс": process,
            "Код": code,
            "Площадка": site,
            "Цеха": ", ".join(sorted(set(data["shops"]))),
            "Птичников на площадке": houses,
            "Средство": agent,
            "Позиция прайса": position,
            "Ед.изм": norm["Ед.изм"],
            "Норма": per_house,
            "Норма на": "птичник",
            "Раствор на 1 птичник, л": solution,
            "Кол-во за месяц": treatments,
            "Расход за месяц": month_qty,
            "Раствор за месяц, л": solution * treatments if solution is not None else None,
            "Цена за ед.": price,
            "Сумма": total,
            "Замечания": "; ".join(problems),
        })

    order = {"ОС": 1, "МС": 2, "ПР": 3, "ДС": 4, "ДЗ": 5, "ГГ": 6, "ГГ+ГФ": 7,
             "ГФ": 8, "ОД": 9, "ПДГ": 10, "СО": 11, "—": 12}
    rows.sort(key=lambda r: (order.get(r["Код"], 99), r["Приложение"], r["Площадка"]))
    return rows, houses_per_shop


def write_sheet(ws, header, data_rows, money_cols=(), number_cols=()):
    ws.append(header)
    for col in range(1, len(header) + 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for values in data_rows:
        ws.append(values)
        r = ws.max_row
        for col in range(1, len(header) + 1):
            ws.cell(r, col).border = BORDER
        for col in money_cols:
            ws.cell(r, col).number_format = "# ##0.00"
        for col in number_cols:
            ws.cell(r, col).number_format = "# ##0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:{}{}".format(chr(64 + len(header)), ws.max_row)


def build_workbook(rows, schedule, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    header = ["Приложение", "Процесс", "Код", "Площадка", "Цеха", "Средство",
              "Позиция прайса", "Птичников на площадке", "Ед.изм",
              "Норма", "Норма на", "Раствор на 1 птичник, л",
              "Кол-во за месяц", "Расход за месяц", "Раствор за месяц, л",
              "Цена за ед.", "Сумма", "Замечания"]
    ws = wb.create_sheet("Расчёт")
    write_sheet(ws, header, [[r[h] for h in header] for r in rows],
                money_cols=(16, 17), number_cols=(10, 12, 14, 15))
    for row_idx, r in enumerate(rows, start=2):
        if r["Замечания"]:
            for col in range(1, len(header) + 1):
                ws.cell(row_idx, col).fill = WARN_FILL
    total_row = ws.max_row + 2
    ws.cell(total_row, 16, "ИТОГО").font = Font(bold=True)
    ws.cell(total_row, 17, "=SUM(Q2:Q{})".format(ws.max_row - 1)).font = Font(bold=True)
    ws.cell(total_row, 17).number_format = "# ##0.00"
    for col, width in zip(range(1, len(header) + 1),
                          (11, 36, 8, 11, 26, 34, 34, 12, 8, 10, 11, 16, 13, 14, 16, 13, 15, 34)):
        ws.column_dimensions[chr(64 + col)].width = width

    # Группируем по позиции прайса: в программе одно и то же средство
    # встречается в разных написаниях («Формалин», «формалин», «Формалин 37,6%»).
    by_agent = defaultdict(lambda: [0.0, 0.0, "", ""])
    for r in rows:
        if r["Расход за месяц"] is None:
            continue
        item = by_agent[r["Позиция прайса"] or r["Средство"]]
        item[0] += r["Расход за месяц"]
        item[1] += r["Сумма"] or 0.0
        item[2] = r["Ед.изм"]
        item[3] = "нет цены" if r["Цена за ед."] is None else ""
    ws2 = wb.create_sheet("Итог по средствам")
    write_sheet(ws2, ["Средство", "Ед.изм", "Расход за месяц", "Сумма", "Замечания"],
                [[k, v[2], v[0], v[1] or None, v[3]] for k, v in sorted(by_agent.items())],
                money_cols=(4,), number_cols=(3,))
    for col, width in zip("ABCDE", (40, 9, 16, 16, 14)):
        ws2.column_dimensions[col].width = width

    by_site = defaultdict(float)
    for r in rows:
        by_site[r["Площадка"]] += r["Сумма"] or 0.0
    ws3 = wb.create_sheet("Итог по площадкам")
    write_sheet(ws3, ["Площадка", "Сумма за месяц"],
                [[k, v or None] for k, v in sorted(by_site.items())], money_cols=(2,))
    for col, width in zip("AB", (16, 18)):
        ws3.column_dimensions[col].width = width

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)


def print_report(rows, schedule):
    print("Период: {} — {}\n".format(schedule["period"]["from"], schedule["period"]["to"]))
    fmt = "{:<38} {:<8} {:<26} {:>10} {:>9} {:>6} {:>7} {:>12} {:>12} {:>14}"
    print(fmt.format("Процесс", "Площадка", "Средство", "норма", "р-р, л",
                     "на", "кол-во", "за месяц", "цена", "сумма"))
    print("-" * 153)
    total = 0.0
    for r in rows:
        total += r["Сумма"] or 0.0
        print(fmt.format(
            r["Процесс"][:38], r["Площадка"], r["Средство"][:26],
            "{:.3f}".format(r["Норма"]) if r["Норма"] is not None else "—",
            "{:g}".format(r["Раствор на 1 птичник, л"]) if r["Раствор на 1 птичник, л"] is not None else "—",
            r["Норма на"][:8],
            r["Кол-во за месяц"],
            "{:.3f}".format(r["Расход за месяц"]) if r["Расход за месяц"] is not None else "—",
            "{:,.2f}".format(r["Цена за ед."]) if r["Цена за ед."] is not None else "нет цены",
            "{:,.2f}".format(r["Сумма"]) if r["Сумма"] is not None else "—"))
    print("-" * 153)
    print("ИТОГО: {:,.2f}".format(total))

    # ОЧН стоит в той же ячейке графика, что и ОС («ОС/ОЧН»), и учтена по ОС.
    covered = {c for r in rows for c in r["Код"].split("+")} | {"ОЧН"}
    in_schedule = {e["code"] for h in schedule["houses"] for e in h["events"]}
    missing = sorted(in_schedule - covered)
    if missing:
        print("\nЕсть в графике, но нормы в санитарной программе не нашлось: "
              + ", ".join(missing))

    if UNIT_MISMATCHES:
        print("\nНорма и цена заданы в разных единицах (считаю 1 л = 1 кг):")
        for agent, (norm_unit, price_unit) in sorted(UNIT_MISMATCHES.items()):
            print("  {}: норма в {}, цена за {}".format(agent[:50], norm_unit, price_unit))

    problems = [r for r in rows if r["Замечания"]]
    if problems:
        print("\nСтроки с замечаниями:")
        for r in problems:
            print("  {} / {} / {}: {}".format(r["Процесс"], r["Площадка"],
                                              r["Средство"][:30], r["Замечания"]))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="график санитарного разрыва")
    ap.add_argument("--sheet", help="лист графика (по умолчанию первый)")
    ap.add_argument("--norms", default="data/normy.csv")
    ap.add_argument("--extra", default="data/normy_dopolnitelno.csv",
                    help="нормы работ, не привязанных к графику санразрыва")
    ap.add_argument("--price", default="data/dezsredstva.csv")
    ap.add_argument("--map", default="data/sopostavlenie.csv")
    ap.add_argument("--out", default="out/raschet.xlsx")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    schedule = parse(args.xlsx, args.sheet)
    norms = read_csv(args.norms)
    prices = load_prices(args.price, args.map)
    rows, houses_per_shop = build_rows(schedule, norms, prices)
    if args.extra and os.path.exists(args.extra):
        rows += build_extra_rows(read_csv(args.extra), prices, houses_per_shop)

    build_workbook(rows, schedule, args.out)
    print("Расчёт сохранён:", args.out)
    if not args.quiet:
        print()
        print_report(rows, schedule)
    return 0


if __name__ == "__main__":
    sys.exit(main())
