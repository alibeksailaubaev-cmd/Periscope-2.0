# -*- coding: utf-8 -*-
"""Разбор санитарной программы АО «УКПФ» в таблицу норм расхода.

В книге санитарной программы приложения №1–17 задают норму расхода на один
птичник по участкам. Скрипт вытаскивает их в единую таблицу и привязывает
каждое приложение к отметке в графике санитарного разрыва.

    python3 src/parse_program.py data/sanitarnaya-programma-31.08.2026.xlsx \
        --out data/normy.csv
"""

import argparse
import csv
import re
import sys

from collections import Counter

import openpyxl

# Участок санитарной программы -> цеха графика санитарного разрыва.
SITE_SHOPS = {
    "А": ["А"],
    "Б": ["Б"],
    "В": ["В"],
    "Г": ["Г 1", "Г 2", "Г 3", "Г 4", "Г 5"],
    "Д": ["Д 1", "Д 2", "Д 3", "Д 4"],
    "Е": ["Е"],
    "Ж": ["Ж"],
    "БП 12": ["БП 12"],
    "БП 13": ["БП 13"],
    "БП 14": ["БП 14"],
}

# Номер участка в заголовке -> буква участка (для строк, где в скобках «БП»).
SITE_BY_NUMBER = {
    "3": "Г", "5": "Д", "7": "Б", "8": "Е", "9": "А", "10": "Ж", "11": "В",
    "12": "БП 12", "13": "БП 13", "14": "БП 14",
}

# Разбор приложений. kind:
#   occurrence — если операция за цикл повторяется (дезинфекция идёт дважды),
#   номер вхождения: 1 — первая по дате, 2 — вторая. Пусто — считать все.
#   «средство» — в колонке value количество самого средства;
#   «раствор»  — в колонке value объём готового раствора, средство = раствор × %.
# В приложениях №2-6 колонка подписана «кол-во дез. раствора», но по факту это
# количество средства (подтверждено заказчиком), а % — концентрация разведения,
# поэтому объём раствора получается делением: средство / концентрация.
# code — отметка в графике санитарного разрыва, по которой считается частота.
APPENDICES = [
    {
        "sheet": "приложение №2,", "no": "2", "rows": (13, 22),
        "process": "Дезинфекция линий поения (санация 2)", "day": "2",
        "code": "МС", "kind": "средство", "solution_from_conc": True,
        "site": "C", "agent": "E", "value": "F", "conc": "G", "unit": "л",
    },
    {
        "sheet": "приложение №3", "no": "3", "rows": (13, 22),
        "process": "Мойка", "day": "2",
        "code": "МС", "kind": "средство", "solution_from_conc": True,
        "site": "C", "agent": "E", "value": "F", "conc": "G", "unit": "л",
    },
    {
        "sheet": "приложение №4", "no": "4", "rows": (13, 22),
        "process": "Мойка", "day": "2",
        "code": "МС", "kind": "средство", "solution_from_conc": True,
        "site": "C", "agent": "E", "value": "F", "conc": "G", "unit": "л",
    },
    {
        "sheet": "приложение №5", "no": "5", "rows": (13, 22),
        "process": "Дезинфекция линий поения (санация 1)", "day": "1",
        "code": "ОС", "kind": "средство", "solution_from_conc": True,
        "site": "C", "agent": "E", "value": "F", "conc": "G", "unit": "кг",
    },
    {
        "sheet": "приложение №6", "no": "6", "rows": (13, 22),
        "process": "Дезинфекция линий поения (санация 3)", "day": "3",
        "code": "ПР", "kind": "средство", "solution_from_conc": True,
        "site": "C", "agent": "E", "value": "F", "conc": "G", "unit": "л",
    },
    {
        "sheet": "приложение №8", "no": "8", "rows": (13, 22),
        "process": "Влажная дезинфекция", "day": "5",
        "code": "ДЗ", "occurrence": 1, "kind": "средство",
        "site": "C", "agent": "E", "value": "F", "conc": "G", "solution": "H", "unit": "кг",
    },
    {
        "sheet": "приложение №9", "no": "9", "rows": (12, 21),
        "process": "Влажная дезинфекция на подстилку", "day": "6",
        "code": "ДЗ", "occurrence": 2, "kind": "средство",
        "site": "C", "agent": "E", "value": "F", "conc": "G", "solution": "H", "unit": "л",
    },
    {
        "sheet": "приложение №11", "no": "11", "rows": (11, 21),
        "process": "Аэрозольная газация формалином", "day": "8",
        "code": "ГФ", "kind": "средство",
        "site": "A", "agent": "C", "value": "D", "conc": "E", "unit": "л",
    },
    {
        "sheet": "приложение №12", "no": "12", "rows": (13, 22),
        "process": "Газация ракет", "day": "7",
        "code": "ГГ", "kind": "средство",
        "site": "A", "agent": "C", "value": "D", "unit": "шт",
    },
    {
        "sheet": "приложение №12", "no": "12", "rows": (13, 22),
        "process": "Газация птичников (подготовка)", "day": "8",
        "code": "ПДГ", "kind": "средство",
        "site": "A", "agent": "C", "value": "E", "unit": "шт",
    },
    {
        "sheet": "приложение №13", "no": "13", "rows": (13, 22),
        "process": "Дезинфекция линий поения (санация 4)", "day": "8",
        "code": "ПДГ", "kind": "средство", "solution_from_conc": True,
        "site": "A", "agent": "C", "value": "D", "conc": "E", "unit": "кг",
    },
    {
        "sheet": "приложение №15", "no": "15", "rows": (11, 15),
        "process": "Обработка септиков", "day": "—",
        "code": "СО", "kind": "средство",
        "site": "C", "agent": "E", "value": "F", "unit": "кг",
    },
    {
        "sheet": "приложение №16", "no": "16", "rows": (11, 21),
        "process": "Аэрозольная газация глутаркой", "day": "7",
        "code": "ГГ", "kind": "средство",
        "site": "A", "agent": "C", "value": "D", "conc": "E", "unit": "л",
    },
    {
        "sheet": "Приложение №17", "no": "17", "rows": (11, 18),
        "process": "Влажная дезинсекция", "day": "4",
        "code": "ДС", "kind": "средство",
        "site": "A", "agent": "B", "value": "C", "conc": "D", "solution": "E", "unit": "л",
    },
]

SITE_RE = re.compile(r"участок\s*№?\s*([\d,\s]+)\s*\(([^)]+)\)", re.IGNORECASE)

# Листы со списками птичников: лист -> ячейка со списком.
HOUSE_LISTS = [("№1 ", "C9"), ("№2", "C9")]
HOUSE_RE = re.compile(r"^(БП\s*\d+|[А-Я])[\s-]*(\d+)?(?:-(\d+))?$")


def count_houses(wb):
    """Цех -> сколько в нём птичников по спискам санитарной программы.

    В части версий программы листов со списками нет — тогда возвращается
    пустой словарь, и число птичников берётся из графика санразрыва.
    """
    counts = Counter()
    for sheet, ref in HOUSE_LISTS:
        if sheet not in wb.sheetnames:
            continue
        text = wb[sheet][ref].value or ""
        for part in re.split(r"[;.]", str(text)):
            part = part.strip()
            if not part.startswith("Птичник"):
                continue
            name = part.replace("Птичник", "").strip()
            if name.startswith("БП"):
                counts["БП " + re.search(r"БП\s*(\d+)", name).group(1)] += 1
                continue
            multi = re.match(r"^([А-Я])\s*(\d+)-(\d+)$", name)
            counts["{} {}".format(multi.group(1), multi.group(2)) if multi
                   else name.split()[0]] += 1
    return counts


def parse_sites(text):
    """'Бройлерный участок №12,13,14 (БП)' -> ['БП 12', 'БП 13', 'БП 14']"""
    if not text:
        return []
    m = SITE_RE.search(str(text))
    if not m:
        return []
    numbers = [n.strip() for n in m.group(1).split(",") if n.strip()]
    letter = m.group(2).strip().upper()
    if letter.startswith("БП"):
        return [SITE_BY_NUMBER[n] for n in numbers if n in SITE_BY_NUMBER]
    return [letter] if letter in SITE_SHOPS else []


def cell(ws, col, row):
    value = ws["{}{}".format(col, row)].value
    return value.strip() if isinstance(value, str) else value


def value_label(ws, spec):
    """Подпись колонки, из которой взята норма — нужна для проверок."""
    for row in range(1, spec["rows"][0]):
        text = cell(ws, spec["value"], row)
        if isinstance(text, str) and text.strip():
            return " ".join(text.split())
    return ""


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    houses = count_houses(wb)
    rows = []
    warnings = []

    for spec in APPENDICES:
        if spec["sheet"] not in wb.sheetnames:
            warnings.append("приложения «{}» в этой версии программы нет — "
                            "процесс «{}» не считается"
                            .format(spec["sheet"], spec["process"]))
            continue
        ws = wb[spec["sheet"]]
        label = value_label(ws, spec)
        seen = set()
        current = None
        for row in range(spec["rows"][0], spec["rows"][1] + 1):
            sites = parse_sites(cell(ws, spec["site"], row))
            agent = cell(ws, spec["agent"], row)
            value = cell(ws, spec["value"], row)
            if not agent or value is None:
                continue
            if sites:
                current = sites
            elif current:
                # Строка-продолжение предыдущего участка (второе значение нормы).
                warnings.append(
                    "прил. №{}: у участка {} несколько норм ({} и {}) — взята первая"
                    .format(spec["no"], ", ".join(current), seen_value(rows, spec, current), value))
                continue
            else:
                continue

            conc = cell(ws, spec["conc"], row) if spec.get("conc") else None
            solution = cell(ws, spec["solution"], row) if spec.get("solution") else None

            if spec["kind"] == "раствор":
                agent_qty = float(value) * float(conc) if conc else None
                solution_qty = float(value)
            else:
                agent_qty = float(value)
                if spec.get("solution_from_conc") and conc:
                    solution_qty = float(value) / float(conc)
                else:
                    solution_qty = float(solution) if solution else None

            for site in sites:
                if site in seen:
                    continue
                seen.add(site)
                for shop in SITE_SHOPS[site]:
                    rows.append({
                        "Приложение": spec["no"],
                        "Процесс": spec["process"],
                        "День санразрыва": spec["day"],
                        "Код в графике": spec["code"],
                        "Вхождение": spec.get("occurrence", ""),
                        "Участок": site,
                        "Цех": shop,
                        "Средство": agent,
                        "Птичников в цехе": houses.get(shop, ""),
                        "Ед.изм": spec["unit"],
                        "Норма средства на птичник": round(agent_qty, 4) if agent_qty is not None else "",
                        "Раствор, л": solution_qty if solution_qty is not None else "",
                        "Концентрация": conc if conc is not None else "",
                        "Как задано": spec["kind"],
                        "Подпись колонки нормы": label,
                    })

    return rows, warnings


def seen_value(rows, spec, sites):
    for r in reversed(rows):
        if r["Приложение"] == spec["no"] and r["Участок"] in sites:
            return r["Норма средства на птичник"]
    return "?"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="книга санитарной программы")
    ap.add_argument("--out", default="data/normy.csv", help="куда сохранить таблицу норм")
    args = ap.parse_args(argv)

    rows, warnings = parse(args.xlsx)
    fields = list(rows[0].keys())
    with open(args.out, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    print("Нормы записаны:", args.out)
    print("  строк:", len(rows))
    print("  приложений:", len({r["Приложение"] for r in rows}))
    houses = {r["Цех"]: r["Птичников в цехе"] for r in rows}
    total = sum(n for n in houses.values() if isinstance(n, int))
    print("  птичников по спискам программы:",
          total if total else "списков нет, возьмутся из графика")
    for w in dict.fromkeys(warnings):
        print("  ВНИМАНИЕ:", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
